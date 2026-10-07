from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from fpdf import FPDF

from app.schemas.transcript import TranscriptResponse

FONT_PATH = Path(__file__).resolve().parent.parent / "assets" / "fonts" / "NotoSansSinhala-Regular.ttf"
FONT_NAME = "NotoSansSinhala"

EXPORT_MEDIA_TYPES = {
    "txt": "text/plain; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


def render_transcript_text(transcript: TranscriptResponse) -> str:
    lines = [transcript.title, ""]
    lines.extend(segment.text for segment in transcript.segments)
    return "\n".join(lines)


def render_transcript_docx(transcript: TranscriptResponse) -> bytes:
    document = Document()
    # Make Sinhala runs use a Sinhala-capable font in Word.
    normal = document.styles["Normal"]
    normal.font.name = "Iskoola Pota"
    normal.element.rPr.rFonts.set(qn("w:cs"), "Iskoola Pota")

    document.add_heading(transcript.title, level=1)
    for segment in transcript.segments:
        document.add_paragraph(segment.text)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_transcript_pdf(transcript: TranscriptResponse) -> bytes:
    pdf = FPDF()
    pdf.add_font(FONT_NAME, "", str(FONT_PATH))
    # HarfBuzz shaping is required for correct Sinhala conjuncts/vowel signs.
    pdf.set_text_shaping(True)
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()

    pdf.set_font(FONT_NAME, size=18)
    pdf.multi_cell(0, 10, transcript.title, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font(FONT_NAME, size=12)
    for segment in transcript.segments:
        pdf.multi_cell(0, 7, segment.text, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    return bytes(pdf.output())
